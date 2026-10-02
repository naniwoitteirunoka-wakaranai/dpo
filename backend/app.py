from flask import Flask, request, jsonify, send_file
from flask_cors import CORS

import sqlite3
import os
import json

from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font

import fitz  # PyMuPDF
import pytesseract
from PIL import Image

from groq import Groq


app = Flask(__name__)
CORS(app)

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
EXPORT_DIR = os.path.join(BASE_DIR, "exports")
DB_NAME = os.path.join(BASE_DIR, "dpo_framework.db")

os.makedirs(EXPORT_DIR, exist_ok=True)


# ==========================================================
# GROQ
# ==========================================================

GROQ_API_KEY = os.environ.get("GROQ_API_KEY")

groq_client = None

if GROQ_API_KEY:
    groq_client = Groq(
        api_key=GROQ_API_KEY
    )


GROQ_MODEL = "openai/gpt-oss-120b"


# ==========================================================
# DATABASE
# ==========================================================

def get_db():

    conn = sqlite3.connect(DB_NAME)

    conn.row_factory = sqlite3.Row

    return conn


def init_db():

    conn = get_db()

    cur = conn.cursor()


    cur.execute("""
        CREATE TABLE IF NOT EXISTS assessments(
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            use_case TEXT NOT NULL,
            description TEXT NOT NULL,
            personal_data INTEGER,
            sensitive_data INTEGER,
            external_ai INTEGER,
            retention INTEGER,
            consent INTEGER,
            affected_users INTEGER,
            risk_score INTEGER,
            risk_level TEXT,
            created_at TEXT
        )
    """)


    cur.execute("""
        CREATE TABLE IF NOT EXISTS incidents(
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            incident_type TEXT NOT NULL,
            description TEXT NOT NULL,
            data_exposed TEXT,
            sensitive_data INTEGER,
            external_ai INTEGER,
            affected_users INTEGER,
            incident_score INTEGER,
            severity TEXT,
            status TEXT DEFAULT 'OPEN',
            created_at TEXT
        )
    """)


    conn.commit()

    conn.close()


init_db()


# ==========================================================
# RISK ENGINE
#
# Formula:
# R = P + 2S + E + D + (1-C)
# ==========================================================

def calculate_risk_score(
    personal,
    sensitive,
    external,
    retention,
    consent
):

    P = int(personal)
    S = int(sensitive)
    E = int(external)
    D = int(retention)
    C = int(consent)

    score = (
        P
        + (2 * S)
        + E
        + D
        + (1 - C)
    )

    return score


def classify_risk(score):

    if score <= 1:
        return "LOW"

    elif score <= 3:
        return "MEDIUM"

    elif score <= 5:
        return "HIGH"

    return "CRITICAL"


def generate_recommendations(
    data,
    level
):

    recs = []


    if data["personal_data"]:

        recs.extend([
            "Apply data minimization before sending prompts to GenAI.",
            "Restrict access to personal information.",
            "Review prompt content for unnecessary identifiers."
        ])


    if data["sensitive_data"]:

        recs.extend([
            "Conduct an additional privacy impact review.",
            "Apply stronger security controls for sensitive information."
        ])


    if data["external_ai"]:

        recs.extend([
            "Review the external GenAI provider before deployment.",
            "Verify data-processing agreements with the provider.",
            "Avoid sending unnecessary personal information."
        ])


    if data["retention"]:

        recs.append(
            "Review retention and deletion policies for GenAI outputs."
        )


    if not data["consent"]:

        recs.append(
            "Require DPO or legal approval before deployment."
        )


    if level in ["HIGH", "CRITICAL"]:

        recs.extend([
            "Perform a formal privacy risk assessment.",
            "Implement monitoring before production deployment."
        ])


    if level == "CRITICAL":

        recs.extend([
            "Escalate the use case to senior privacy governance.",
            "Block deployment until mitigation controls are completed."
        ])


    return list(
        dict.fromkeys(recs)
    )


# ==========================================================
# INCIDENT ENGINE
# ==========================================================

FINANCIAL_KEYWORDS = [
    "account",
    "bank",
    "upi",
    "card",
    "credit",
    "debit",
    "financial"
]


def calculate_incident_score(
    sensitive,
    external,
    users,
    exposed_text
):

    score = 0


    if sensitive:

        score += 2


    if external:

        score += 1


    if users >= 500:

        score += 2

    elif users >= 100:

        score += 1


    text = exposed_text.lower()


    if any(
        word in text
        for word in FINANCIAL_KEYWORDS
    ):

        score += 1


    return score


def classify_incident(score):

    if score <= 1:

        return "LOW"

    elif score <= 3:

        return "MEDIUM"

    elif score <= 5:

        return "HIGH"

    return "CRITICAL"


def incident_actions(
    severity,
    data
):

    actions = [
        "Contain the incident immediately.",
        "Identify exposed personal data.",
        "Estimate affected individuals.",
        "Document the incident for DPO review."
    ]


    if data["sensitive_data"]:

        actions.append(
            "Assess impact of sensitive data exposure."
        )


    if data["external_ai"]:

        actions.append(
            "Review external AI provider interaction and logs."
        )


    if severity in ["HIGH", "CRITICAL"]:

        actions.extend([
            "Escalate incident for DPO review.",
            "Review notification and escalation requirements.",
            "Begin remediation activities."
        ])


    if severity == "CRITICAL":

        actions.extend([
            "Pause affected GenAI workflow until investigation completes.",
            "Perform root-cause analysis.",
            "Implement additional privacy safeguards."
        ])


    actions.append(
        "Close incident after remediation is verified."
    )


    return list(
        dict.fromkeys(actions)
    )


# ==========================================================
# DOCUMENT EXTRACTION
# ==========================================================

ALLOWED_EXTENSIONS = {
    ".pdf",
    ".txt"
}


def extract_pdf_text(file_path):

    document = fitz.open(
        file_path
    )

    extracted_pages = []


    for page in document:

        text = page.get_text(
            "text"
        ).strip()


        if len(text) >= 30:

            extracted_pages.append(
                text
            )

            continue


        # --------------------------------------------------
        # OCR fallback
        # --------------------------------------------------

        pixmap = page.get_pixmap(
            matrix=fitz.Matrix(
                2,
                2
            ),
            alpha=False
        )


        image = Image.frombytes(
            "RGB",
            [
                pixmap.width,
                pixmap.height
            ],
            pixmap.samples
        )


        ocr_text = pytesseract.image_to_string(
            image
        ).strip()


        if ocr_text:

            extracted_pages.append(
                ocr_text
            )


    document.close()


    return "\n\n".join(
        extracted_pages
    )


def extract_document_text(file):

    filename = file.filename or ""


    extension = os.path.splitext(
        filename
    )[1].lower()


    if extension not in ALLOWED_EXTENSIONS:

        raise ValueError(
            "Unsupported file type. Please upload a PDF or TXT file."
        )


    temp_path = os.path.join(
        EXPORT_DIR,
        (
            f"_temp_"
            f"{datetime.now().strftime('%Y%m%d%H%M%S%f')}"
            f"{extension}"
        )
    )


    file.save(
        temp_path
    )


    try:

        if extension == ".pdf":

            text = extract_pdf_text(
                temp_path
            )


        elif extension == ".txt":

            with open(
                temp_path,
                "r",
                encoding="utf-8",
                errors="ignore"
            ) as f:

                text = f.read()


        else:

            text = ""


    finally:

        if os.path.exists(
            temp_path
        ):

            os.remove(
                temp_path
            )


    if not text.strip():

        raise ValueError(
            "Could not extract any text from the uploaded document."
        )


    return text.strip()


# ==========================================================
# AI RISK EXTRACTION
# ==========================================================

def extract_risk_parameters_with_ai(
    document_text
):

    if not groq_client:

        raise RuntimeError(
            "GROQ_API_KEY is not configured on the server."
        )


    system_prompt = """

You are a privacy-risk extraction assistant for a GenAI DPO
Privacy Framework.

Your job is NOT to calculate the risk score.

Your job is to read an incident report, privacy report,
security report, or other uploaded document and extract
structured information required by the framework.

Extract ONLY information supported by the document.

If a value cannot be determined from the document:

- boolean fields should be false
- affected_users should be 0
- use_case should be a concise description inferred from
  the document
- description should summarize the relevant situation

Definitions:

personal_data:
true if the document indicates that information relating
to identifiable individuals is processed, exposed, uploaded,
stored, shared, or otherwise handled.

sensitive_data:
true if the document indicates sensitive/special-category
personal information or similarly high-risk personal data.

external_ai:
true if an external Generative AI / AI provider or third-party
AI service is involved.

retention:
true if the document indicates that the AI provider,
organization, system, logs, prompts, outputs, or related data
are retained or stored.

consent:
true ONLY when the document clearly indicates that valid
consent or another explicitly documented lawful basis exists.
If the document does not establish this, return false.

affected_users:
Return the number of affected individuals when explicitly
stated or reasonably extractable from the document.
Otherwise return 0.

Important:
Do not invent facts.
Do not estimate affected users.
Do not calculate risk_score.
Do not calculate risk_level.
"""


    user_prompt = f"""

Analyze the following document and extract the privacy-risk
parameters required by the framework.

DOCUMENT:

{document_text}

"""


    response = groq_client.chat.completions.create(

        model=GROQ_MODEL,

        messages=[
            {
                "role": "system",
                "content": system_prompt
            },
            {
                "role": "user",
                "content": user_prompt
            }
        ],

        temperature=0,

        response_format={
            "type": "json_schema",

            "json_schema": {

                "name":
                    "genai_privacy_assessment",

                "schema": {

                    "type":
                        "object",

                    "properties": {

                        "use_case": {
                            "type":
                                "string"
                        },

                        "description": {
                            "type":
                                "string"
                        },

                        "personal_data": {
                            "type":
                                "boolean"
                        },

                        "sensitive_data": {
                            "type":
                                "boolean"
                        },

                        "external_ai": {
                            "type":
                                "boolean"
                        },

                        "retention": {
                            "type":
                                "boolean"
                        },

                        "consent": {
                            "type":
                                "boolean"
                        },

                        "affected_users": {
                            "type":
                                "integer"
                        }

                    },

                    "required": [
                        "use_case",
                        "description",
                        "personal_data",
                        "sensitive_data",
                        "external_ai",
                        "retention",
                        "consent",
                        "affected_users"
                    ],

                    "additionalProperties":
                        False
                }
            }
        }
    )


    content = response.choices[0].message.content


    return json.loads(
        content
    )


# ==========================================================
# AI INCIDENT EXTRACTION
# ==========================================================

def extract_incident_parameters_with_ai(
    document_text
):

    if not groq_client:

        raise RuntimeError(
            "GROQ_API_KEY is not configured on the server."
        )


    system_prompt = """

You are an incident-analysis extraction assistant for a
GenAI DPO Privacy Framework.

Your job is NOT to calculate the incident score.

Your job is to read an incident report, privacy report,
security report, breach report, or similar document and
extract ONLY the incident parameters required by the
framework.

Extract ONLY information supported by the document.

If a value cannot be determined:

- sensitive_data should be false
- external_ai should be false
- affected_users should be 0
- incident_type should be a concise incident category
- description should summarize the incident
- data_exposed should describe the exposed data when known,
  otherwise return "Unknown"

Definitions:

incident_type:
A concise category describing what happened.
Examples include data exposure, unauthorized disclosure,
AI-related data breach, privacy violation, credential exposure,
or security incident.

description:
A concise factual summary of the incident.

data_exposed:
Describe the personal, sensitive, financial, credential,
system, or other information exposed or potentially exposed.
Do not invent data that is not supported by the document.

sensitive_data:
true if the document indicates sensitive personal data,
special-category data, health data, financial data, identity
information, child-related sensitive information, or similarly
high-risk personal information was exposed or involved.

external_ai:
true if an external Generative AI provider, AI service,
chatbot, model provider, or third-party AI system is involved.

affected_users:
Return the number of affected individuals when explicitly
stated or reasonably extractable from the document.
Otherwise return 0.

Important:

Do not calculate incident_score.

Do not calculate severity.

Do not invent facts.

Do not estimate affected users.

Only extract facts supported by the document.
"""


    user_prompt = f"""

Analyze the following incident document and extract the
incident-management parameters required by the framework.

DOCUMENT:

{document_text}

"""


    response = groq_client.chat.completions.create(

        model=GROQ_MODEL,

        messages=[
            {
                "role": "system",
                "content": system_prompt
            },
            {
                "role": "user",
                "content": user_prompt
            }
        ],

        temperature=0,

        response_format={
            "type": "json_schema",

            "json_schema": {

                "name":
                    "genai_incident_assessment",

                "schema": {

                    "type":
                        "object",

                    "properties": {

                        "incident_type": {
                            "type":
                                "string"
                        },

                        "description": {
                            "type":
                                "string"
                        },

                        "data_exposed": {
                            "type":
                                "string"
                        },

                        "sensitive_data": {
                            "type":
                                "boolean"
                        },

                        "external_ai": {
                            "type":
                                "boolean"
                        },

                        "affected_users": {
                            "type":
                                "integer"
                        }

                    },

                    "required": [
                        "incident_type",
                        "description",
                        "data_exposed",
                        "sensitive_data",
                        "external_ai",
                        "affected_users"
                    ],

                    "additionalProperties":
                        False
                }
            }
        }
    )


    content =
        response.choices[0].message.content


    return json.loads(
        content
    )


# ==========================================================
# HELPERS
# ==========================================================

def validation_error(message):

    return jsonify({
        "success": False,
        "error": message
    }), 400


def current_time():

    return datetime.now().strftime(
        "%Y-%m-%d %H:%M:%S"
    )


def serialize_rows(rows):

    return [
        dict(row)
        for row in rows
    ]


# ==========================================================
# HEALTH API
# ==========================================================

@app.route(
    "/api/health",
    methods=["GET"]
)
def health():

    return jsonify({
        "status": "ok",
        "service":
            "GenAI DPO Framework API"
    })


# ==========================================================
# RISK ASSESSMENT API
# ==========================================================

@app.route(
    "/api/assess",
    methods=["POST"]
)
def assess():

    data =
        request.get_json(
            silent=True
        )


    if not data:

        return validation_error(
            "Invalid JSON body."
        )


    required = [
        "use_case",
        "description",
        "personal_data",
        "sensitive_data",
        "external_ai",
        "retention",
        "consent"
    ]


    for field in required:

        if field not in data:

            return validation_error(
                f"Missing field: {field}"
            )


    score =
        calculate_risk_score(
            data["personal_data"],
            data["sensitive_data"],
            data["external_ai"],
            data["retention"],
            data["consent"]
        )


    level =
        classify_risk(
            score
        )


    recommendations =
        generate_recommendations(
            data,
            level
        )


    conn =
        get_db()

    cur =
        conn.cursor()


    cur.execute("""
        INSERT INTO assessments(
            use_case,
            description,
            personal_data,
            sensitive_data,
            external_ai,
            retention,
            consent,
            affected_users,
            risk_score,
            risk_level,
            created_at
        )
        VALUES(?,?,?,?,?,?,?,?,?,?,?)
    """, (

        data["use_case"],

        data["description"],

        int(
            data["personal_data"]
        ),

        int(
            data["sensitive_data"]
        ),

        int(
            data["external_ai"]
        ),

        int(
            data["retention"]
        ),

        int(
            data["consent"]
        ),

        data.get(
            "affected_users",
            0
        ),

        score,

        level,

        current_time()
    ))


    conn.commit()


    assessment_id =
        cur.lastrowid


    conn.close()


    return jsonify({

        "success":
            True,

        "assessment_id":
            assessment_id,

        "risk_score":
            score,

        "risk_level":
            level,

        "recommendations":
            recommendations
    })


# ==========================================================
# AI DOCUMENT RISK ASSESSMENT API
# ==========================================================

@app.route(
    "/api/assess-ai",
    methods=["POST"]
)
def assess_ai():

    if "file" not in request.files:

        return validation_error(
            "No document uploaded."
        )


    file =
        request.files["file"]


    if not file.filename:

        return validation_error(
            "No file selected."
        )


    try:

        document_text =
            extract_document_text(
                file
            )


        ai_data =
            extract_risk_parameters_with_ai(
                document_text
            )


        ai_data["personal_data"] =
            bool(
                ai_data["personal_data"]
            )


        ai_data["sensitive_data"] =
            bool(
                ai_data["sensitive_data"]
            )


        ai_data["external_ai"] =
            bool(
                ai_data["external_ai"]
            )


        ai_data["retention"] =
            bool(
                ai_data["retention"]
            )


        ai_data["consent"] =
            bool(
                ai_data["consent"]
            )


        ai_data["affected_users"] =
            max(
                0,
                int(
                    ai_data.get(
                        "affected_users",
                        0
                    )
                )
            )


        score =
            calculate_risk_score(
                ai_data["personal_data"],
                ai_data["sensitive_data"],
                ai_data["external_ai"],
                ai_data["retention"],
                ai_data["consent"]
            )


        level =
            classify_risk(
                score
            )


        recommendations =
            generate_recommendations(
                ai_data,
                level
            )


        conn =
            get_db()

        cur =
            conn.cursor()


        cur.execute("""
            INSERT INTO assessments(
                use_case,
                description,
                personal_data,
                sensitive_data,
                external_ai,
                retention,
                consent,
                affected_users,
                risk_score,
                risk_level,
                created_at
            )
            VALUES(?,?,?,?,?,?,?,?,?,?,?)
        """, (

            ai_data["use_case"],

            ai_data["description"],

            int(
                ai_data["personal_data"]
            ),

            int(
                ai_data["sensitive_data"]
            ),

            int(
                ai_data["external_ai"]
            ),

            int(
                ai_data["retention"]
            ),

            int(
                ai_data["consent"]
            ),

            ai_data["affected_users"],

            score,

            level,

            current_time()
        ))


        conn.commit()


        assessment_id =
            cur.lastrowid


        conn.close()


        return jsonify({

            "success":
                True,

            "assessment_id":
                assessment_id,

            "extracted_data": {

                "use_case":
                    ai_data["use_case"],

                "description":
                    ai_data["description"],

                "personal_data":
                    ai_data["personal_data"],

                "sensitive_data":
                    ai_data["sensitive_data"],

                "external_ai":
                    ai_data["external_ai"],

                "retention":
                    ai_data["retention"],

                "consent":
                    ai_data["consent"],

                "affected_users":
                    ai_data["affected_users"]
            },

            "risk_score":
                score,

            "risk_level":
                level,

            "recommendations":
                recommendations
        })


    except ValueError as error:

        return validation_error(
            str(error)
        )


    except Exception as error:

        print(
            "AI assessment error:",
            str(error)
        )


        return jsonify({

            "success":
                False,

            "error":
                "Unable to analyze the uploaded document."
        }), 500


# ==========================================================
# ASSESSMENT HISTORY
# ==========================================================

@app.route(
    "/api/assessments",
    methods=["GET"]
)
def assessments():

    conn =
        get_db()


    rows =
        conn.execute("""
            SELECT *
            FROM assessments
            ORDER BY id DESC
        """).fetchall()


    conn.close()


    return jsonify({

        "success":
            True,

        "count":
            len(rows),

        "data":
            serialize_rows(
                rows
            )
    })


# ==========================================================
# INCIDENT MANAGEMENT API
# ==========================================================

@app.route(
    "/api/incidents",
    methods=["POST"]
)
def incidents():

    data =
        request.get_json(
            silent=True
        )


    if not data:

        return validation_error(
            "Invalid JSON body."
        )


    required = [
        "incident_type",
        "description",
        "data_exposed",
        "sensitive_data",
        "external_ai",
        "affected_users"
    ]


    for field in required:

        if field not in data:

            return validation_error(
                f"Missing field: {field}"
            )


    users =
        int(
            data["affected_users"]
        )


    if users < 0:

        return validation_error(
            "Affected users cannot be negative."
        )


    score =
        calculate_incident_score(
            data["sensitive_data"],
            data["external_ai"],
            users,
            data["data_exposed"]
        )


    severity =
        classify_incident(
            score
        )


    actions =
        incident_actions(
            severity,
            data
        )


    conn =
        get_db()

    cur =
        conn.cursor()


    cur.execute("""
        INSERT INTO incidents(
            incident_type,
            description,
            data_exposed,
            sensitive_data,
            external_ai,
            affected_users,
            incident_score,
            severity,
            status,
            created_at
        )
        VALUES(?,?,?,?,?,?,?,?,?,?)
    """, (

        data["incident_type"],

        data["description"],

        data["data_exposed"],

        int(
            data["sensitive_data"]
        ),

        int(
            data["external_ai"]
        ),

        users,

        score,

        severity,

        "OPEN",

        current_time()
    ))


    conn.commit()


    incident_id =
        cur.lastrowid


    conn.close()


    return jsonify({

        "success":
            True,

        "incident_id":
            incident_id,

        "incident_score":
            score,

        "severity":
            severity,

        "actions":
            actions
    })


# ==========================================================
# AI INCIDENT ASSESSMENT API
# ==========================================================

@app.route(
    "/api/incident-ai",
    methods=["POST"]
)
def incident_ai():

    # ------------------------------------------------------
    # Validate uploaded document
    # ------------------------------------------------------

    if "file" not in request.files:

        return validation_error(
            "No document uploaded."
        )


    file =
        request.files["file"]


    if not file.filename:

        return validation_error(
            "No file selected."
        )


    try:

        # --------------------------------------------------
        # Extract document text
        # --------------------------------------------------

        document_text =
            extract_document_text(
                file
            )


        # --------------------------------------------------
        # AI extracts incident parameters
        # --------------------------------------------------

        ai_data =
            extract_incident_parameters_with_ai(
                document_text
            )


        # --------------------------------------------------
        # Normalize extracted values
        # --------------------------------------------------

        ai_data["incident_type"] =
            str(
                ai_data.get(
                    "incident_type",
                    "GenAI Privacy Incident"
                )
            ).strip()


        ai_data["description"] =
            str(
                ai_data.get(
                    "description",
                    ""
                )
            ).strip()


        ai_data["data_exposed"] =
            str(
                ai_data.get(
                    "data_exposed",
                    "Unknown"
                )
            ).strip()


        ai_data["sensitive_data"] =
            bool(
                ai_data.get(
                    "sensitive_data",
                    False
                )
            )


        ai_data["external_ai"] =
            bool(
                ai_data.get(
                    "external_ai",
                    False
                )
            )


        ai_data["affected_users"] =
            max(
                0,
                int(
                    ai_data.get(
                        "affected_users",
                        0
                    )
                )
            )


        # --------------------------------------------------
        # IMPORTANT:
        #
        # AI DOES NOT calculate the incident score.
        #
        # The deterministic Python incident engine does.
        # --------------------------------------------------

        score =
            calculate_incident_score(
                ai_data["sensitive_data"],
                ai_data["external_ai"],
                ai_data["affected_users"],
                ai_data["data_exposed"]
            )


        severity =
            classify_incident(
                score
            )


        actions =
            incident_actions(
                severity,
                ai_data
            )


        # --------------------------------------------------
        # Store AI-generated incident exactly like
        # a manually assessed incident.
        # --------------------------------------------------

        conn =
            get_db()

        cur =
            conn.cursor()


        cur.execute("""
            INSERT INTO incidents(
                incident_type,
                description,
                data_exposed,
                sensitive_data,
                external_ai,
                affected_users,
                incident_score,
                severity,
                status,
                created_at
            )
            VALUES(?,?,?,?,?,?,?,?,?,?)
        """, (

            ai_data["incident_type"],

            ai_data["description"],

            ai_data["data_exposed"],

            int(
                ai_data["sensitive_data"]
            ),

            int(
                ai_data["external_ai"]
            ),

            ai_data["affected_users"],

            score,

            severity,

            "OPEN",

            current_time()
        ))


        conn.commit()


        incident_id =
            cur.lastrowid


        conn.close()


        # --------------------------------------------------
        # Return extracted parameters to frontend.
        # --------------------------------------------------

        return jsonify({

            "success":
                True,

            "incident_id":
                incident_id,

            "extracted_data": {

                "incident_type":
                    ai_data["incident_type"],

                "description":
                    ai_data["description"],

                "data_exposed":
                    ai_data["data_exposed"],

                "sensitive_data":
                    ai_data["sensitive_data"],

                "external_ai":
                    ai_data["external_ai"],

                "affected_users":
                    ai_data["affected_users"]
            },

            "incident_score":
                score,

            "severity":
                severity,

            "actions":
                actions
        })


    except ValueError as error:

        return validation_error(
            str(error)
        )


    except Exception as error:

        print(
            "AI incident assessment error:",
            str(error)
        )


        return jsonify({

            "success":
                False,

            "error":
                "Unable to analyze the incident document."
        }), 500


# ==========================================================
# INCIDENT HISTORY
# ==========================================================

@app.route(
    "/api/incidents",
    methods=["GET"]
)
def get_incidents():

    conn =
        get_db()


    rows =
        conn.execute("""
            SELECT *
            FROM incidents
            ORDER BY id DESC
        """).fetchall()


    conn.close()


    return jsonify({

        "success":
            True,

        "count":
            len(rows),

        "data":
            serialize_rows(
                rows
            )
    })


# ==========================================================
# DASHBOARD STATS
# ==========================================================

@app.route(
    "/api/stats",
    methods=["GET"]
)
def stats():

    conn =
        get_db()


    total_assessments =
        conn.execute(
            "SELECT COUNT(*) FROM assessments"
        ).fetchone()[0]


    high_assessments =
        conn.execute("""
            SELECT COUNT(*)
            FROM assessments
            WHERE risk_level IN ('HIGH','CRITICAL')
        """).fetchone()[0]


    open_incidents =
        conn.execute("""
            SELECT COUNT(*)
            FROM incidents
            WHERE status='OPEN'
        """).fetchone()[0]


    critical_incidents =
        conn.execute("""
            SELECT COUNT(*)
            FROM incidents
            WHERE severity='CRITICAL'
        """).fetchone()[0]


    conn.close()


    return jsonify({

        "success":
            True,

        "total_assessments":
            total_assessments,

        "high_critical_assessments":
            high_assessments,

        "open_incidents":
            open_incidents,

        "critical_incidents":
            critical_incidents
    })


# ==========================================================
# EXCEL EXPORT
# ==========================================================

@app.route(
    "/api/export",
    methods=["GET"]
)
def export_excel():

    conn =
        get_db()


    assessments =
        conn.execute("""
            SELECT *
            FROM assessments
            ORDER BY id DESC
        """).fetchall()


    incidents =
        conn.execute("""
            SELECT *
            FROM incidents
            ORDER BY id DESC
        """).fetchall()


    conn.close()


    wb =
        Workbook()


    # ------------------------------------------------------
    # Risk Assessments
    # ------------------------------------------------------

    ws =
        wb.active

    ws.title =
        "Risk Assessments"


    headers = [
        "ID",
        "Use Case",
        "Description",
        "Personal Data",
        "Sensitive Data",
        "External AI",
        "Retention",
        "Consent",
        "Affected Users",
        "Risk Score",
        "Risk Level",
        "Created At"
    ]


    for col, header in enumerate(
        headers,
        start=1
    ):

        cell =
            ws.cell(
                row=1,
                column=col
            )

        cell.value =
            header

        cell.font =
            Font(
                bold=True
            )


    for item in assessments:

        ws.append([

            item["id"],

            item["use_case"],

            item["description"],

            "YES"
            if item["personal_data"]
            else "NO",

            "YES"
            if item["sensitive_data"]
            else "NO",

            "YES"
            if item["external_ai"]
            else "NO",

            "YES"
            if item["retention"]
            else "NO",

            "YES"
            if item["consent"]
            else "NO",

            item["affected_users"],

            item["risk_score"],

            item["risk_level"],

            item["created_at"]
        ])


    # ------------------------------------------------------
    # Incidents
    # ------------------------------------------------------

    ws2 =
        wb.create_sheet(
            "Incidents"
        )


    headers2 = [
        "ID",
        "Incident Type",
        "Description",
        "Sensitive Data",
        "External AI",
        "Affected Users",
        "Incident Score",
        "Severity",
        "Status",
        "Created At"
    ]


    for col, header in enumerate(
        headers2,
        start=1
    ):

        cell =
            ws2.cell(
                row=1,
                column=col
            )

        cell.value =
            header

        cell.font =
            Font(
                bold=True
            )


    for item in incidents:

        ws2.append([

            item["id"],

            item["incident_type"],

            item["description"],

            "YES"
            if item["sensitive_data"]
            else "NO",

            "YES"
            if item["external_ai"]
            else "NO",

            item["affected_users"],

            item["incident_score"],

            item["severity"],

            item["status"],

            item["created_at"]
        ])


    filename = os.path.join(

        EXPORT_DIR,

        (
            "GenAI_DPO_Assessments_"
            f"{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        )
    )


    wb.save(
        filename
    )


    return send_file(

        filename,

        as_attachment=True,

        download_name=
            "GenAI_DPO_Framework_Report.xlsx"
    )


# ==========================================================
# ROOT API
# ==========================================================

@app.route(
    "/",
    methods=["GET"]
)
def root():

    return jsonify({

        "message":
            "GenAI DPO Framework Backend Running",

        "health":
            "/api/health"
    })


# ==========================================================
# APP RUNNER
# ==========================================================

if __name__ == "__main__":

    port =
        int(
            os.environ.get(
                "PORT",
                5000
            )
        )


    app.run(

        host="0.0.0.0",

        port=port,

        debug=False
    )
